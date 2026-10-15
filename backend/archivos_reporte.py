"""
Armado de archivos de reporte: PDF (reportlab) y Excel (openpyxl).

Son funciones puras (datos -> bytes) y este módulo importa sólo lo que necesitan: server.py las
corre en un proceso aparte (_generar_archivo) para que un reporte de miles de filas no frene las
demás peticiones.
"""
import io
from copy import copy
from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

PARAGUAY_TZ = ZoneInfo("America/Asuncion")


def now_paraguay():
    return datetime.now(PARAGUAY_TZ)


ESTILO_TABLA_REPORTE = [
    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0044CC')),
    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ('FONTSIZE', (0, 0), (-1, 0), 10),
    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
    ('BACKGROUND', (0, 1), (-1, -1), colors.white),
    ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
    ('FONTSIZE', (0, 1), (-1, -1), 9),
    ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ('ALIGN', (-1, 1), (-1, -1), 'RIGHT'),
    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
]
ESTILO_TOTALES_REPORTE = [
    ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E0E0E0')),
    ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
]
BANDAS_REPORTE = [colors.white, colors.HexColor('#F5F5F5')]


class _TablaPorTramos(Flowable):
    """Tabla del reporte que reportlab arma de a tramos de filas.

    Una Table con decenas de miles de filas se vuelve a crear con todas las que quedan en cada
    salto de página (50.000 filas: ~90 s con el proceso ocupado). Así cada página trabaja con
    un tramo corto. El PDF sale igual que con una sola tabla: encabezado en cada página, bandas
    que arrancan en blanco en cada página, mismos anchos de columna y fila de totales al final."""
    TRAMO = 100  # filas por tramo; siempre más de las que entran en una página (~40)

    def __init__(self, encabezado, filas, totales, col_widths, desde=0):
        Flowable.__init__(self)
        self.encabezado, self.filas, self.totales = encabezado, filas, totales
        self.col_widths, self.desde = col_widths, desde
        self.hAlign = 'CENTER'
        self._tabla = None

    def _armar(self):
        hasta = self.desde + self.TRAMO
        ultimo = hasta >= len(self.filas)
        datos = [self.encabezado] + self.filas[self.desde:hasta]
        if ultimo and self.totales:
            datos.append(self.totales)
        tabla = Table(datos, repeatRows=1, colWidths=self.col_widths)
        tabla.setStyle(TableStyle(ESTILO_TABLA_REPORTE + [('ROWBACKGROUNDS', (0, 1), (-1, -1), BANDAS_REPORTE)]))
        if ultimo and self.totales:
            tabla.setStyle(TableStyle(ESTILO_TOTALES_REPORTE))
        return tabla, ultimo

    def wrap(self, availWidth, availHeight):
        self._tabla, self._ultimo = self._armar()
        ancho, alto = self._tabla.wrap(availWidth, availHeight)
        if not self._ultimo:
            alto = max(alto, availHeight + 1)  # quedan filas en otros tramos: que se parta
        self.width, self.height = ancho, alto
        return ancho, alto

    def split(self, availWidth, availHeight):
        # El frame llama a wrap y después a split: se reutiliza la tabla ya armada
        if self._tabla is None:
            self.wrap(availWidth, availHeight)
        tabla, ultimo = self._tabla, self._ultimo
        partes = tabla.split(availWidth, availHeight)
        if ultimo or not partes:
            return partes
        entraron = len(partes[0]._cellvalues) - 1  # filas de datos en esta página (sin encabezado)
        if entraron <= 0:
            return []
        return [partes[0], _TablaPorTramos(self.encabezado, self.filas, self.totales,
                                           self.col_widths, self.desde + entraron)]

    def draw(self):
        self._tabla.drawOn(self.canv, 0, 0)


def crear_pdf_reporte(titulo, subtitulo, columnas, datos, totales=None, col_widths=None):
    """Genera un PDF con tabla de datos"""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    # Title
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=18, spaceAfter=6, alignment=TA_CENTER)
    elements.append(Paragraph(titulo, title_style))

    # Subtitle
    subtitle_style = ParagraphStyle('CustomSubtitle', parent=styles['Normal'], fontSize=10, spaceAfter=20, alignment=TA_CENTER, textColor=colors.grey)
    elements.append(Paragraph(subtitulo, subtitle_style))

    # Table: sin anchos dados, reportlab los calcula sobre todas las filas; se calculan una vez
    # con la tabla completa y los tramos los reutilizan
    if col_widths is None:
        completa = Table([columnas] + datos + ([totales] if totales else []), repeatRows=1)
        completa.setStyle(TableStyle(ESTILO_TABLA_REPORTE + (ESTILO_TOTALES_REPORTE if totales else [])))
        completa._calc_width(doc.width)
        col_widths = completa._colWidths
    elements.append(_TablaPorTramos(columnas, datos, totales, col_widths))

    # Footer
    elements.append(Spacer(1, 20))
    footer_style = ParagraphStyle('Footer', parent=styles['Normal'], fontSize=8, textColor=colors.grey, alignment=TA_CENTER)
    elements.append(Paragraph(f"Generado el {now_paraguay().strftime('%d/%m/%Y %H:%M')} - Ethercloud ERP", footer_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

def crear_excel_reporte(titulo, subtitulo, columnas, datos, totales=None):
    """Genera un Excel (.xlsx) con tabla de datos estilizada"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte"

    num_cols = len(columnas)
    last_col = get_column_letter(num_cols)

    # Title row (row 1)
    ws.merge_cells(f"A1:{last_col}1")
    title_cell = ws["A1"]
    title_cell.value = titulo
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="0044CC")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Subtitle row (row 2)
    ws.merge_cells(f"A2:{last_col}2")
    sub_cell = ws["A2"]
    sub_cell.value = subtitulo
    sub_cell.font = Font(name="Calibri", size=10, color="808080")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Empty row 3 (spacer)
    ws.row_dimensions[3].height = 6

    # Headers (row 4)
    HEADER_ROW = 4
    header_fill = PatternFill(start_color="0044CC", end_color="0044CC", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, col_name in enumerate(columnas, 1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[HEADER_ROW].height = 20

    # Data rows. Las dos primeras filas se estilan celda por celda; las demás copian el estilo
    # ya registrado de la fila con su misma banda (asignar Font/Fill/Alignment nuevos en cada
    # celda obligaba a openpyxl a buscarlos en su tabla de estilos: 50.000 filas, ~30 s).
    data_font = Font(name="Calibri", size=9)
    for row_offset, row_data in enumerate(datos, 1):
        row_num = HEADER_ROW + row_offset
        if row_offset <= 2:
            bg = "FFFFFF" if row_offset % 2 == 1 else "F5F5F5"
            fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_idx, value=value)
                cell.fill = fill
                cell.font = data_font
                cell.border = border
                cell.alignment = Alignment(horizontal="left" if col_idx < num_cols else "right", vertical="center")
        else:
            modelo = HEADER_ROW + (1 if row_offset % 2 == 1 else 2)
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_idx, value=value)
                cell._style = copy(ws.cell(row=modelo, column=col_idx)._style)

    # Totals row
    if totales:
        tot_row = HEADER_ROW + len(datos) + 1
        tot_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
        tot_font = Font(name="Calibri", size=9, bold=True)
        for col_idx, value in enumerate(totales, 1):
            cell = ws.cell(row=tot_row, column=col_idx, value=value)
            cell.fill = tot_fill
            cell.font = tot_font
            cell.border = border
            cell.alignment = Alignment(horizontal="right", vertical="center")

    # Auto-size columns
    for col_idx in range(1, num_cols + 1):
        max_len = len(str(columnas[col_idx - 1]))
        for row_data in datos:
            if col_idx <= len(row_data):
                max_len = max(max_len, len(str(row_data[col_idx - 1])))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def crear_excel_ranking_bi(productos, metrica, metrica_label, periodo_label, titulo, subtitulo):
    """Excel del ranking de productos del BI, con gráfico de los 20 primeros."""
    # Build workbook manually (with chart)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ranking Productos"

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = titulo
    t.font = Font(name="Calibri", size=14, bold=True, color="0044CC")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Subtitle
    ws.merge_cells("A2:H2")
    s = ws["A2"]
    s.value = subtitulo
    s.font = Font(name="Calibri", size=9, color="808080")
    s.alignment = Alignment(horizontal="center")

    # Headers row 4
    headers = ["#", "Producto", "Ventas (₲)", "Ganancia (₲)", "Unidades",
               f"Prom. uds/{periodo_label}", "Stock actual", "Variación %"]
    hfill = PatternFill(start_color="0044CC", end_color="0044CC", fill_type="solid")
    hfont = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=ci, value=h)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[4].height = 20

    # Data rows: las dos primeras se estilan celda por celda y las demás copian el estilo de la
    # fila de su misma banda (estilos nuevos en cada celda: ~45.000 productos, más de 1 minuto)
    for ri, p in enumerate(productos, 1):
        rn = 4 + ri
        row_vals = [ri, p["nombre"], p["ventas"], p["ganancia"], p["unidades"],
                    p["prom"], p["stock"], p["variacion"]]
        if ri <= 2:
            bg = "FFFFFF" if ri % 2 == 1 else "F5F7FF"
            fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            for ci, val in enumerate(row_vals, 1):
                cell = ws.cell(row=rn, column=ci, value=val)
                cell.fill = fill
                cell.font = Font(name="Calibri", size=9)
                cell.border = border
                cell.alignment = Alignment(horizontal="right" if ci > 1 else "center", vertical="center")
        else:
            modelo = 4 + (1 if ri % 2 == 1 else 2)
            for ci, val in enumerate(row_vals, 1):
                cell = ws.cell(row=rn, column=ci, value=val)
                cell._style = copy(ws.cell(row=modelo, column=ci)._style)

    # Column widths
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 35
    for col in ["C", "D", "E", "F", "G", "H"]:
        ws.column_dimensions[col].width = 16

    # Totals row
    tot_rn = 4 + len(productos) + 1
    tot_fill = PatternFill(start_color="E8EDF5", end_color="E8EDF5", fill_type="solid")
    tot_font = Font(name="Calibri", size=9, bold=True)
    totals = [
        "", "TOTAL",
        sum(p["ventas"] for p in productos),
        sum(p["ganancia"] for p in productos),
        sum(p["unidades"] for p in productos),
        "", "", ""
    ]
    for ci, val in enumerate(totals, 1):
        cell = ws.cell(row=tot_rn, column=ci, value=val)
        cell.fill = tot_fill
        cell.font = tot_font
        cell.border = border
        cell.alignment = Alignment(horizontal="right", vertical="center")

    # Embedded BarChart (top 20)
    chart_data_count = min(20, len(productos))
    chart_col = {"ventas": 3, "ganancia": 4, "unidades": 5}[metrica]
    chart = BarChart()
    chart.type = "bar"  # horizontal
    chart.title = f"Top {chart_data_count} — {metrica_label}"
    chart.y_axis.title = "Producto"
    chart.x_axis.title = metrica_label
    chart.style = 10
    chart.width = 22
    chart.height = 14

    data_ref = Reference(ws, min_col=chart_col, min_row=4, max_row=4 + chart_data_count)
    cats_ref = Reference(ws, min_col=2, min_row=5, max_row=4 + chart_data_count)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.series[0].graphicalProperties.solidFill = "0044CC"

    # Place chart to the right of data (column J, row 4)
    ws.add_chart(chart, "J4")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
